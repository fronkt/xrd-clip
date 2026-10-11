"""Build the URTC 2026 talk deck (paper ID 69, xrd-clip) as .pptx for Google Slides import."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "PA69.pptx")

NAVY = RGBColor(0x14, 0x24, 0x3D)
BLUE = RGBColor(0x1F, 0x77, 0xB4)    # matches XRD-CLIP line in the figures
RED = RGBColor(0xD6, 0x27, 0x28)
GREEN = RGBColor(0x2C, 0xA0, 0x2C)
GREY = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xEE, 0xF3, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x29, 0x37)
FONT = "Arial"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
TOTAL = 11


def text(slide, x, y, w, h, runs, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """runs: str, or list of paragraphs; each paragraph is str or list of (str, {opts})."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        segs = [(para, {})] if isinstance(para, str) else para
        for s, o in segs:
            r = p.add_run()
            r.text = s
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=20, gap=10):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        segs = [(item, {})] if isinstance(item, str) else item
        r0 = p.add_run()
        r0.text = "■  "
        r0.font.name, r0.font.size, r0.font.color.rgb = FONT, Pt(size * 0.6), BLUE
        for s, o in segs:
            r = p.add_run()
            r.text = s
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.color.rgb = o.get("color", INK)
    return tb


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    return s


def label_box(slide, x, y, w, h, fill, lines, size=16, color=WHITE, line=None):
    s = box(slide, x, y, w, h, fill, line)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, para in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        segs = [(para, {})] if isinstance(para, str) else para
        for t, o in segs:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.color.rgb = o.get("color", color)
    return s


def arrow(slide, x, y, w, h, color=GREY):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def new_slide(title, n, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    if kicker:
        text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.4), kicker, size=14,
             color=BLUE, bold=True)
    text(s, Inches(0.6), Inches(0.65), Inches(12.2), Inches(0.9), title, size=32,
         color=NAVY, bold=True)
    text(s, Inches(0.6), Inches(7.0), Inches(9), Inches(0.35),
         "Cai · Polymorph-Resolving Retrieval from PXRD · IEEE MIT URTC 2026",
         size=11, color=GREY)
    text(s, Inches(11.9), Inches(7.0), Inches(0.9), Inches(0.35), f"{n} / {TOTAL}",
         size=11, color=GREY, align=PP_ALIGN.RIGHT)
    return s


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# 1 ── Title ────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(1.2), Inches(0.09))
acc.fill.solid(); acc.fill.fore_color.rgb = BLUE; acc.line.fill.background()
text(s, Inches(0.8), Inches(1.55), Inches(11.8), Inches(2.6),
     ["Polymorph-Resolving Retrieval",
      "of Molecular Crystals from",
      "Powder X-Ray Diffraction"], size=40, color=WHITE, bold=True, spacing=1.0)
text(s, Inches(0.8), Inches(4.4), Inches(11.8), Inches(0.6),
     "A contrastive pattern↔structure model that tells same-molecule crystal forms apart",
     size=20, color=RGBColor(0xC7, 0xD6, 0xEA))
text(s, Inches(0.8), Inches(5.4), Inches(11.5), Inches(1.2),
     [[("Frank Cai", {"bold": True, "size": 22})],
      "Purdue University, West Lafayette, IN  ·  frankyc11223@gmail.com"],
     size=16, color=WHITE)
text(s, Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.4),
     "2026 IEEE MIT Undergraduate Research Technology Conference  ·  Paper ID 69  ·  Track 4A",
     size=13, color=RGBColor(0x9F, 0xB3, 0xCE))
notes(s, "[~20 s] Hi, I'm Frank Cai from Purdue. This talk is about identifying which crystal "
         "form of a molecule you have, straight from its powder X-ray diffraction pattern, "
         "including the hard case where the candidates are the same molecule.")

# 2 ── Problem ──────────────────────────────────────────────────────────────
s = new_slide("Same molecule, different crystal: polymorphs", 2, "THE PROBLEM")
bullets(s, Inches(0.6), Inches(1.75), Inches(6.2), Inches(3.9), [
    [("Polymorphs", {"bold": True}), (" are different packings of the same molecule", {})],
    "Same chemical formula, often the same space group",
    "Only the packing differs, which shows up as shifts in PXRD peak positions and intensities",
    [("It matters: ", {"bold": True}),
     ("ritonavir (an HIV drug) had to be pulled when a less soluble form appeared in production", {})],
], size=20, gap=14)
label_box(s, Inches(0.6), Inches(5.55), Inches(6.1), Inches(1.15), NAVY,
          [[("ROY: same formula, same space group (P2₁/n).", {"bold": True})],
           "Only the PXRD pattern tells Y and OP apart."], size=15)
# Real example: ROY yellow (Y) vs orange-plate (OP) forms, PyMOL outline renders + CCDC-simulated PXRD
fig_h = Inches(5.25)
s.shapes.add_picture(os.path.join(HERE, "polymorphs", "roy_Y_vs_OP.png"),
                     Inches(6.95), Inches(1.5), height=fig_h)
notes(s, "[~50 s] Polymorphs are the same molecule packed differently. This is ROY, a molecule "
         "with many polymorphs named by their colour. The yellow and orange-plate forms have the "
         "same formula and the same space group; the molecule just twists differently, 105 versus "
         "46 degrees, and packs differently. Their simulated powder patterns, at the bottom, are "
         "clearly different. That is the only signal that separates them. This matters in "
         "practice: ritonavir's marketed form was withdrawn when a less soluble polymorph "
         "appeared in manufacturing. (Structures: CSD QAXMEH01 and QAXMEH03.)")

# 3 ── Gap / task ───────────────────────────────────────────────────────────
s = new_slide("Existing models lean on composition and symmetry", 3, "THE GAP")
bullets(s, Inches(0.6), Inches(1.75), Inches(6.3), Inches(3.5), [
    [("Classical search-match: ", {"bold": True}), ("correlate the pattern against a reference library", {})],
    [("XCCP", {"bold": True}), (" (155k inorganic crystals) and ", {}), ("XtalNet", {"bold": True}),
     (" (MOFs): contrastive PXRD↔structure retrieval", {})],
    "Both lean on composition and space-group cues, and that signal disappears for polymorphs",
    "No prior work tests retrieval with the composition held fixed",
], size=20, gap=14)
label_box(s, Inches(7.3), Inches(1.8), Inches(5.4), Inches(0.6), BLUE,
          [[("Our task: within-family retrieval", {"bold": True})]], size=18)
label_box(s, Inches(7.3), Inches(2.5), Inches(5.4), Inches(2.7), LIGHT,
          ["Query: one polymorph's PXRD pattern",
           "Library: only its sibling polymorphs",
           "Pick the packing that produced it", "",
           [("A composition/symmetry retriever is at chance ", {"color": INK}),
            ("by construction", {"bold": True, "color": INK})]],
          size=17, color=INK)
label_box(s, Inches(7.3), Inches(5.4), Inches(5.4), Inches(1.15), NAVY,
          [[("Chance = 0.380", {"bold": True, "size": 26})],
           "(average over family sizes)"], size=14)
notes(s, "[~50 s] The usual tool is search-match: correlate against a library. Recent "
         "contrastive models such as XCCP and XtalNet do learned retrieval, but on inorganic "
         "crystals and MOFs, and they lean on composition and symmetry; XCCP reports 93% "
         "space-group ID. For polymorphs those cues are gone. So the task here is "
         "within-family retrieval: the query is one polymorph's pattern, the library is only "
         "its siblings. Anything that uses only formula or symmetry scores 0.380, the "
         "family-size chance rate.")

# 4 ── Method ───────────────────────────────────────────────────────────────
s = new_slide("Method: a CLIP-style dual encoder", 4, "APPROACH")
y0 = Inches(1.95)
label_box(s, Inches(0.6), y0, Inches(2.3), Inches(1.15), LIGHT,
          [[("PXRD pattern", {"bold": True})], "5–50° 2θ, 2,251 pts"], size=15, color=INK)
arrow(s, Inches(3.0), y0 + Inches(0.4), Inches(0.55), Inches(0.35))
label_box(s, Inches(3.65), y0, Inches(2.9), Inches(1.15), BLUE,
          [[("1D ResNet", {"bold": True})], "wide kernels + global pooling"], size=15)
arrow(s, Inches(6.65), y0 + Inches(0.4), Inches(0.55), Inches(0.35))
y1 = Inches(3.75)
label_box(s, Inches(0.6), y1, Inches(2.3), Inches(1.15), LIGHT,
          [[("Crystal structure", {"bold": True})], "atoms + unit cell"], size=15, color=INK)
arrow(s, Inches(3.0), y1 + Inches(0.4), Inches(0.55), Inches(0.35))
label_box(s, Inches(3.65), y1, Inches(2.9), Inches(1.15), RED,
          [[("Atom-set Transformer", {"bold": True})], "periodic coords + lattice token"], size=15)
arrow(s, Inches(6.65), y1 + Inches(0.4), Inches(0.55), Inches(0.35))
label_box(s, Inches(7.3), y0, Inches(2.4), Inches(2.9), NAVY,
          [[("Shared 256-d", {"bold": True})], [("embedding", {"bold": True})], "",
           "match = cosine", "similarity"], size=16)
arrow(s, Inches(9.8), Inches(3.15), Inches(0.55), Inches(0.35))
label_box(s, Inches(10.45), y0, Inches(2.3), Inches(2.9), LIGHT,
          [[("Symmetric InfoNCE", {"bold": True})], "",
           "pull matched pairs together", "", "push the other B−1 apart"], size=15, color=INK)
bullets(s, Inches(0.6), Inches(5.35), Inches(12.2), Inches(1.5), [
    "Train on 247,089 CSD molecular crystals; patterns are re-simulated on the fly, so no pattern is seen twice",
    "Retrieval: embed the query pattern, rank library structures by cosine similarity",
], size=18, gap=8)
notes(s, "[~50 s] Two encoders map into one embedding space. The pattern side is a 1D residual "
         "CNN: a diffractogram is a 1D signal where local peak shape carries the physics. The "
         "structure side is a permutation-invariant Transformer over atoms with a lattice "
         "token, because peak positions are set by the unit cell. Training uses the CLIP "
         "symmetric InfoNCE loss: each pattern should match its own structure and not the "
         "other structures in the batch. At test time we rank structures by cosine similarity.")

# 5 ── Training details ─────────────────────────────────────────────────────
s = new_slide("Physics-based augmentation, polymorph-aware training", 5, "APPROACH")
label_box(s, Inches(0.6), Inches(1.8), Inches(6.0), Inches(0.55), BLUE,
          [[("Randomized measurement model (every epoch)", {"bold": True})]], size=17)
bullets(s, Inches(0.7), Inches(2.5), Inches(5.9), Inches(3.6), [
    "2θ zero offset (sample displacement)",
    "Isotropic lattice strain (thermal / compositional)",
    "Caglioti + Scherrer peak broadening",
    "March–Dollase preferred orientation",
    "Polynomial + amorphous-hump background",
    "Poisson counting noise",
], size=18, gap=8)
label_box(s, Inches(7.0), Inches(1.8), Inches(5.75), Inches(0.55), RED,
          [[("Making the model resolve polymorphs", {"bold": True})]], size=17)
bullets(s, Inches(7.1), Inches(2.5), Inches(5.6), Inches(3.6), [
    [("Hard negatives: ", {"bold": True}), ("sibling polymorphs share a batch, so the loss must use packing, not formula", {})],
    [("Family-safe split: ", {"bold": True}), ("whole families stay in one split, on top of a scaffold split", {})],
    [("Verified: 0 of 10,795", {"bold": True}), (" families cross a split boundary", {})],
], size=18, gap=12)
label_box(s, Inches(0.6), Inches(5.75), Inches(12.15), Inches(0.95), LIGHT,
          ["CSD: 315,060 molecular crystals → 10,795 polymorph families  ·  "
           "split 247,089 / 30,887 / 30,886  ·  test: 1,077 families, 2,835 targets  ·  60 epochs"],
          size=15, color=INK)
notes(s, "[~45 s] Two design choices. First, augmentation: every training pattern is "
         "re-simulated through a randomized instrument model with displacement, strain, "
         "broadening, texture, background and noise, so the query looks like a real "
         "measurement and not like the clean library entry. Second, polymorph-aware "
         "training: sibling polymorphs are put in the same batch as hard negatives, and whole "
         "families go into a single split, so no sibling of a test structure is ever seen in "
         "training. I checked: zero of 10,795 families leak across splits.")

# 6 ── Headline ─────────────────────────────────────────────────────────────
s = new_slide("The model resolves polymorphs well above chance", 6, "RESULT")
for i, (val, lab, col) in enumerate([("0.561", "XRD-CLIP\nwithin-family top-1", BLUE),
                                      ("0.380", "composition-only\nchance", GREY),
                                      ("1.48×", "lift over\nchance", NAVY)]):
    x = Inches(0.6 + i * 4.1)
    box(s, x, Inches(1.9), Inches(3.85), Inches(2.5), LIGHT)
    text(s, x, Inches(2.0), Inches(3.85), Inches(1.3), val, size=60, bold=True, color=col,
         align=PP_ALIGN.CENTER)
    text(s, x, Inches(3.3), Inches(3.85), Inches(0.9), lab.split("\n"), size=17, color=INK,
         align=PP_ALIGN.CENTER)
text(s, Inches(0.6), Inches(4.65), Inches(12.2), Inches(0.5),
     "Top-2 0.843  ·  Top-5 0.949  ·  MRR 0.738   (1,077 held-out families, 2,835 queries)",
     size=19, color=INK, bold=True, align=PP_ALIGN.CENTER)
label_box(s, Inches(0.6), Inches(5.4), Inches(12.15), Inches(1.2), NAVY,
          [[("Honest context: ", {"bold": True}),
            ("on a perfectly clean query, plain correlation scores 1.0 because the query is a byte-identical copy "
             "of the library entry. The real test is a query that carries measurement effects the library doesn't.", {})]],
          size=16)
notes(s, "[~40 s] On clean patterns the model picks the right polymorph 56% of the time "
         "against a 38% chance floor, a 1.48x lift, and it's in the top two 84% of the time. "
         "One caveat up front: on a perfectly clean query, plain correlation gets 100% because "
         "the query is literally the library entry. That's not a realistic setting. What "
         "matters is what happens when the measured pattern carries real instrument effects.")

# 7 ── Sweeps ───────────────────────────────────────────────────────────────
s = new_slide("Robust to the effects that break search-match", 7, "ROBUSTNESS")
s.shapes.add_picture(os.path.join(HERE, "fig_sweeps.png"), Inches(0.6), Inches(1.6), width=Inches(12.15))
bullets(s, Inches(0.6), Inches(5.85), Inches(12.2), Inches(1.1), [
    [("Peak shifts (a, b): ", {"bold": True}),
     ("the model stays flat (0.561 → 0.523) while correlation collapses (1.0 → 0.298)", {})],
    [("Limitation (c): ", {"bold": True}),
     ("severe broadening at 10 nm crystallites, where correlation (0.884) beats the model (0.472)", {})],
], size=17, gap=6)
notes(s, "[~55 s] Here we degrade the query one dial at a time; the library stays clean. Blue "
         "is our model, red is full-profile correlation, green is peak-list matching, dashed "
         "is chance. Under sample displacement and lattice strain, which are the shifts that "
         "break search-match in real labs, the model is flat while correlation drops below "
         "chance. Panel c is the honest limitation: severe Scherrer broadening keeps peak "
         "positions, and correlation handles that better than we do.")

# 8 ── Sign-cancellation control ────────────────────────────────────────────
s = new_slide("When shifts add, classical methods fall to chance", 8, "ROBUSTNESS")
s.shapes.add_picture(os.path.join(HERE, "fig_combined.png"), Inches(0.6), Inches(1.55), width=Inches(12.15))
bullets(s, Inches(0.6), Inches(5.8), Inches(12.2), Inches(1.15), [
    "Displacement and lattice expansion shift peaks in opposite directions. Same sizes, relative sign flipped:",
    [("Correlation 0.807 → 0.334, peak-list 0.784 → 0.340 (chance); ", {"bold": True}),
     ("model 0.555 → 0.534", {"bold": True, "color": BLUE})],
], size=17, gap=6)
notes(s, "[~55 s] With everything applied at once (panel a), classical methods look decent. "
         "But that's partly luck: sample displacement and thermal expansion shift peaks in "
         "opposite directions, so they can cancel. Panel b is the control: same shift sizes, "
         "only the relative sign flipped. When the shifts add, correlation and peak-list both "
         "drop to chance, while the model barely moves. In a real sample you don't know in "
         "advance which case you're in, and roughly half the time it's the bad one. The "
         "model's invariance doesn't depend on the sign.")

# 9 ── Ablations ────────────────────────────────────────────────────────────
s = new_slide("Ablations: where the robustness comes from", 9, "ANALYSIS")
rows = [("Within-family top-1", "Full aug (main)", "No shift/strain aug", "No aug (clean)"),
        ("Clean", "0.561", "0.553", "0.557"),
        ("Pure shift @ 0.1°", "0.552", "0.549", "0.552"),
        ("Field (mild)", "0.554", "0.560", "0.380"),
        ("Field (harsh)", "0.521", "0.528", "0.379"),
        ("Shifts add", "0.534", "0.530", "0.379")]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.75), Inches(7.4), Inches(3.6)).table
for c, wid in enumerate([2.6, 1.6, 1.6, 1.6]):
    tbl.columns[c].width = Inches(wid)
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = v
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        f = p.runs[0].font
        f.name, f.size = FONT, Pt(15 if r else 14)
        f.bold = r == 0 or (c == 3 and r >= 3)
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = NAVY; f.color.rgb = WHITE
        elif c == 3 and r >= 3:
            cell.fill.fore_color.rgb = RGBColor(0xFD, 0xE2, 0xE2); f.color.rgb = RED
        else:
            cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT; f.color.rgb = INK
text(s, Inches(0.6), Inches(5.45), Inches(7.4), Inches(0.4),
     "Same architecture, data, and 60-epoch schedule; only the augmentation changes. Chance = 0.380.",
     size=13, color=GREY)
label_box(s, Inches(8.4), Inches(1.75), Inches(4.35), Inches(1.95), BLUE,
          [[("Peak-shift invariance is EMERGENT", {"bold": True})], "",
           "Dropping shift/strain augmentation changes nothing: the CNN with global pooling is built to tolerate shifts"],
          size=15)
label_box(s, Inches(8.4), Inches(3.9), Inches(4.35), Inches(2.25), RED,
          [[("Profile robustness is LEARNED", {"bold": True})], "",
           "Trained on clean patterns, the twin drops to chance once real instrument effects appear, so this is not fingerprint leakage"],
          size=15)
notes(s, "[~55 s] Two ablations answer where this comes from. Remove only the shift and strain "
         "augmentation: nothing changes, so shift invariance comes from the architecture. "
         "Remove all augmentation: the model still works on clean and pure-shift queries, but "
         "falls to exactly chance once broadening, texture, noise and background appear. That "
         "answers the obvious objection, 'isn't this just memorizing simulated fingerprints?' "
         "If it were, the clean-trained twin would still work. It doesn't, so the "
         "augmentation is teaching real generalization.")

# 10 ── Real data + limitations ─────────────────────────────────────────────
s = new_slide("Sim-to-real spot check and limitations", 10, "SCOPE")
label_box(s, Inches(0.6), Inches(1.8), Inches(6.0), Inches(0.55), BLUE,
          [[("Real experimental patterns (Cu-Kα₁)", {"bold": True})]], size=17)
bullets(s, Inches(0.7), Inches(2.5), Inches(5.9), Inches(3.8), [
    [("Racemic ibuprofen (IBPRAC04): ", {"bold": True}),
     ("true structure ranks 1,872 of 30,887 (top 6.1%) from a real measured pattern the model never saw", {})],
    [("Phenobarbital Forms I/II: ", {"bold": True}),
     ("the determination is not in the library, so no rank applies; no plausible false match surfaced", {})],
    "Small N: a spot check, not a benchmark",
], size=17, gap=12)
label_box(s, Inches(7.0), Inches(1.8), Inches(5.75), Inches(0.55), RED,
          [[("Limitations → next steps", {"bold": True})]], size=17)
bullets(s, Inches(7.1), Inches(2.5), Inches(5.6), Inches(3.8), [
    "Evaluation is mostly on the physics simulator → scale to tens of real molecular patterns",
    "Concedes severe broadening to correlation",
    "Top-1 0.561 leaves headroom → low/high-angle encoder split, Rietveld reranking",
    "Molecular (CSD) only; inorganic and MOFs are already covered",
], size=17, gap=12)
notes(s, "[~40 s] As a first look at transfer, I took real published powder patterns. For "
         "ibuprofen, a real pattern the model never saw puts the true structure in the top 6% "
         "of a 30,000-structure library, which is good but not top-1. This is a spot check, "
         "not a benchmark, and most of the evaluation is still on the simulator. The model "
         "also concedes heavy broadening, and 56% top-1 leaves room to improve.")

# 11 ── Takeaways ───────────────────────────────────────────────────────────
s = new_slide("Takeaways", 11, "CONCLUSION")
for i, (head, body, col) in enumerate([
    ("New task", "Polymorph-resolving PXRD retrieval on CSD molecular crystals, where composition/symmetry methods are at chance by construction", BLUE),
    ("Robust result", "0.561 vs 0.380 chance, and it holds under peak shifts that send classical search-match to chance, whatever their sign", NAVY),
    ("Understood mechanism", "Shift invariance is architectural; instrument-profile robustness is learned from physics-based augmentation, not leakage", RED)]):
    x = Inches(0.6 + i * 4.1)
    box(s, x, Inches(1.85), Inches(3.85), Inches(0.12), col, shape=MSO_SHAPE.RECTANGLE)
    text(s, x, Inches(2.1), Inches(3.85), Inches(0.6), head, size=22, bold=True, color=col)
    text(s, x, Inches(2.75), Inches(3.85), Inches(2.4), body, size=17, color=INK)
label_box(s, Inches(0.6), Inches(5.3), Inches(12.15), Inches(1.35), NAVY,
          [[("Thank you. Questions?", {"bold": True, "size": 26})],
           "Frank Cai · frankyc11223@gmail.com"], size=16)
notes(s, "[~25 s] To sum up: a new polymorph retrieval task where formula and symmetry don't "
         "help; a model that resolves polymorphs above chance and stays robust where classical "
         "search-match fails; and ablations that show why. Next, the same retriever could rank "
         "candidates for PXRD-conditioned crystal structure generation. Thank you, I'm happy "
         "to take questions.")

prs.save(OUT)
print("saved", OUT)
